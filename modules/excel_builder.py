import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def format_number_br(num):
    """Formata número com separadores de milhar no padrão brasileiro"""
    if isinstance(num, int):
        return f"{num:,}".replace(',', '.')
    return str(num) if num is not None else ""


def build_final_workbook(output_filename, playlist_title, tracks, artists_data):
    """
    Constrói o arquivo Excel (.xlsx) completo com as 4 abas estruturadas
    e formatadas de acordo com o modelo original do usuário.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove sheet padrão em branco

    # Estilos padronizados
    header_font = Font(name='Aptos Narrow', size=11, bold=True, color='000000')
    data_font = Font(name='Aptos Narrow', size=11, bold=False, color='000000')
    link_font = Font(name='Aptos Narrow', size=11, bold=False, color='0000FF', underline='single')
    gray_font = Font(name='Aptos Narrow', size=11, bold=False, color='595959')

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )

    # -------------------------------------------------------------
    # ABA 1: Faixas da Playlist
    # -------------------------------------------------------------
    if 'top 50' in playlist_title.lower():
        clean_name = playlist_title.lower().replace('top 50 -', '').replace('top 50', '').strip()
        sheet1_name = f"Top 50 - {clean_name[:20].title()}" if clean_name else "Top 50"
    else:
        sheet1_name = f"Playlist - {playlist_title[:20]}"

    ws1 = wb.create_sheet(title=sheet1_name[:31])
    headers1 = ['', 'Song', 'Artist', 'Genres', 'Album', 'Spotify Track Id', 'ISRC']
    ws1.append(headers1)

    for t in tracks:
        for artist in t.get('artists', []):
            ws1.append([
                t.get('position', ''),
                f" {t.get('name', '')}",
                artist.get('name', ''),
                t.get('genres', ''),
                t.get('album', ''),
                t.get('track_id', ''),
                t.get('isrc', '')
            ])

    # -------------------------------------------------------------
    # ABA 2: Todos os artistas
    # -------------------------------------------------------------
    count_header = 'Músicas\nno top 50' if 'top 50' in playlist_title.lower() else 'Músicas\nna playlist'
    ws2 = wb.create_sheet(title="Todos os artistas")
    headers2 = [
        count_header, 'Artista', 'Ouvintes mensais\ndo Spotify',
        'IG', 'Tiktok', 'Youtube',
        'Número', 'Prioridade', '', 'Marca', '', 'Situação', 'Titular', 'Classe', 'Procurador'
    ]
    ws2.append(headers2)

    # -------------------------------------------------------------
    # ABA 3: Artistas sem marca (Leads de prospecção)
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Artistas sem marca")
    headers3 = [count_header, 'Artista', 'Ouvintes mensais\ndo Spotify', 'IG', 'Tiktok', 'Youtube']
    ws3.append(headers3)

    # -------------------------------------------------------------
    # ABA 4: Artistas com marca
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Artistas com marca")
    headers4 = ['Músicas\nno top 50', 'Artista', 'Número', 'Prioridade', '', 'Marca', '', 'Situação', 'Titular', 'Classe']
    ws4.append(headers4)

    # -------------------------------------------------------------
    # Preenchimento e Formatação das Linhas
    # -------------------------------------------------------------
    for name, data in artists_data.items():
        count_disp = data.get('count_display', '')
        listeners_val = data.get('monthly_listeners')
        listeners_str = format_number_br(listeners_val) if listeners_val else "N/D"
        spotify_url = data.get('spotify_url')

        ig_display = data.get('ig_display', 'N/D')
        ig_url = data.get('ig_url')

        tt_display = data.get('tiktok_display', 'N/D')
        tt_url = data.get('tiktok_url')

        yt_display = data.get('youtube_display', 'N/D')
        yt_url = data.get('youtube_url')

        inpi_records = data.get('inpi_records')
        has_brand = bool(inpi_records and len(inpi_records) > 0)
        inpi_direct_link = data.get('inpi_search_url')

        if has_brand:
            # Artista possui marca(s) registrada(s) ou processo ativo
            for idx, rec in enumerate(inpi_records):
                # Linha da Aba 2 (Todos os artistas)
                row2 = [
                    count_disp if idx == 0 else "",
                    name,  # Sempre preenche o nome do artista em todas as linhas
                    listeners_str if idx == 0 else "",
                    ig_display if idx == 0 else "",
                    tt_display if idx == 0 else "",
                    yt_display if idx == 0 else "",
                    str(rec.get('numero', '')),
                    str(rec.get('prioridade', '')),
                    "",
                    str(rec.get('marca', '')),
                    "",
                    str(rec.get('situacao', '')),
                    str(rec.get('titular', '')),
                    str(rec.get('classe', '')),
                    str(rec.get('procurador', ''))
                ]
                ws2.append(row2)
                r2_idx = ws2.max_row

                if idx == 0:
                    if spotify_url and listeners_val:
                        ws2.cell(r2_idx, 3).hyperlink = spotify_url
                        ws2.cell(r2_idx, 3).font = link_font
                    if ig_url and ig_display != 'N/D':
                        ws2.cell(r2_idx, 4).hyperlink = ig_url
                        ws2.cell(r2_idx, 4).font = link_font
                    if tt_url and tt_display != 'N/D':
                        ws2.cell(r2_idx, 5).hyperlink = tt_url
                        ws2.cell(r2_idx, 5).font = link_font
                    if yt_url and yt_display != 'N/D':
                        ws2.cell(r2_idx, 6).hyperlink = yt_url
                        ws2.cell(r2_idx, 6).font = link_font

                if rec.get('numero_link'):
                    ws2.cell(r2_idx, 7).hyperlink = rec['numero_link']
                    ws2.cell(r2_idx, 7).font = link_font

                # Linha da Aba 4 (Artistas com marca)
                row4 = [
                    count_disp if idx == 0 else "",
                    name,  # Sempre preenche o nome do artista em todas as linhas
                    str(rec.get('numero', '')),
                    str(rec.get('prioridade', '')),
                    "",
                    str(rec.get('marca', '')),
                    "",
                    str(rec.get('situacao', '')),
                    str(rec.get('titular', '')),
                    str(rec.get('classe', ''))
                ]
                ws4.append(row4)
                r4_idx = ws4.max_row
                if rec.get('numero_link'):
                    ws4.cell(r4_idx, 3).hyperlink = rec['numero_link']
                    ws4.cell(r4_idx, 3).font = link_font

        else:
            # Artista sem marca cadastrada (Leads Comerciais)
            row2 = [
                count_disp,
                name,
                listeners_str,
                ig_display,
                tt_display,
                yt_display,
                "Pesquisar INPI" if inpi_direct_link else "N/A",
                "N/A", "", "N/A", "", "N/A", "N/A", "N/A", ""
            ]
            ws2.append(row2)
            r2_idx = ws2.max_row

            if spotify_url and listeners_val:
                ws2.cell(r2_idx, 3).hyperlink = spotify_url
                ws2.cell(r2_idx, 3).font = link_font
            if ig_url and ig_display != 'N/D':
                ws2.cell(r2_idx, 4).hyperlink = ig_url
                ws2.cell(r2_idx, 4).font = link_font
            if tt_url and tt_display != 'N/D':
                ws2.cell(r2_idx, 5).hyperlink = tt_url
                ws2.cell(r2_idx, 5).font = link_font
            if yt_url and yt_display != 'N/D':
                ws2.cell(r2_idx, 6).hyperlink = yt_url
                ws2.cell(r2_idx, 6).font = link_font
            if inpi_direct_link:
                ws2.cell(r2_idx, 7).hyperlink = inpi_direct_link
                ws2.cell(r2_idx, 7).font = link_font

            # Linha da Aba 3 (Artistas sem marca)
            row3 = [count_disp, name, listeners_str, ig_display, tt_display, yt_display]
            ws3.append(row3)
            r3_idx = ws3.max_row

            if spotify_url and listeners_val:
                ws3.cell(r3_idx, 3).hyperlink = spotify_url
                ws3.cell(r3_idx, 3).font = link_font
            if ig_url and ig_display != 'N/D':
                ws3.cell(r3_idx, 4).hyperlink = ig_url
                ws3.cell(r3_idx, 4).font = link_font
            if tt_url and tt_display != 'N/D':
                ws3.cell(r3_idx, 5).hyperlink = tt_url
                ws3.cell(r3_idx, 5).font = link_font
            if yt_url and yt_display != 'N/D':
                ws3.cell(r3_idx, 6).hyperlink = yt_url
                ws3.cell(r3_idx, 6).font = link_font

    # -------------------------------------------------------------
    # Estilização e Ajuste Fino das Abas
    # -------------------------------------------------------------
    for ws in [ws1, ws2, ws3, ws4]:
        ws.row_dimensions[1].height = 26
        # Estilo cabeçalho
        for cell in ws[1]:
            cell.font = header_font
            cell.alignment = align_center

        # Ajuste de largura automática das colunas
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if '\n' in val:
                    val = max(val.split('\n'), key=len)
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_filename)
    return output_filename
